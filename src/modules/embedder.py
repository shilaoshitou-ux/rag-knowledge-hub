"""Parse uploaded documents and persist deterministic Chroma indexes."""

from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Sequence

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from docx import Document as WordDocument

from .model_config import COMPANY_BASE_URL, DEFAULT_EMBEDDING_MODEL, LOCAL_OLLAMA_BASE_URL, validate_base_url

EMBEDDING_MODEL = DEFAULT_EMBEDDING_MODEL
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
MAX_FILE_MB = 100
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
ONLINE_BATCH_SIZE = 8  # Keep each /embeddings request small for proxy limits.


class EmbeddingBatchError(RuntimeError):
    """Safe progress metadata; never include document text or request payloads."""

    def __init__(self, start: int, end: int, total: int, status_code: int | None):
        self.status_code = status_code
        self.start = start
        self.end = end
        self.total = total
        status = f"（HTTP {status_code}）" if status_code else ""
        super().__init__(f"第 {start}–{end} / {total} 个文本片段的向量化或入库失败{status}。")


@dataclass(frozen=True)
class UploadedDocument:
    name: str
    content: bytes


@dataclass
class IndexInfo:
    vectorstore: Chroma
    collection_id: str
    file_count: int
    chunk_count: int
    cached: bool


def display_name(name: str) -> str:
    """Do not trust path fragments provided in an upload's filename."""
    return Path(name.replace("\\", "/")).name


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("gb18030")


def parse_document(file: UploadedDocument) -> list[Document]:
    name = display_name(file.name)
    if not name or not file.content:
        raise ValueError(f"{name or '文件'}为空。")
    if len(file.content) > MAX_FILE_BYTES:
        raise ValueError(f"{name}超过 {MAX_FILE_MB} MB 限制。")

    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(file.content))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(Document(page_content=text, metadata={
                    "source": name, "location": f"第 {page_number} 页", "page": page_number,
                }))
        documents = pages
    elif suffix == ".docx":
        word = WordDocument(BytesIO(file.content))
        documents = [Document(page_content=paragraph.text, metadata={
            "source": name, "location": f"第 {number} 段", "paragraph": number,
        }) for number, paragraph in enumerate(word.paragraphs, 1) if paragraph.text.strip()]
        documents.extend(Document(page_content=" | ".join(cell.text for cell in row.cells), metadata={
            "source": name, "location": f"表格 {table_number} 第 {row_number} 行",
        }) for table_number, table in enumerate(word.tables, 1)
            for row_number, row in enumerate(table.rows, 1))
    elif suffix in {".md", ".markdown"}:
        documents = [Document(page_content=_decode_text(file.content), metadata={
            "source": name, "location": "全文",
        })]
    elif suffix == ".txt":
        documents = [Document(page_content=_decode_text(file.content), metadata={
            "source": name, "location": "全文",
        })]
    elif suffix == ".csv":
        reader = csv.DictReader(StringIO(_decode_text(file.content)))
        if not reader.fieldnames:
            raise ValueError(f"{name}缺少 CSV 表头。")
        documents = [
            Document(
                page_content="\n".join(f"{key}: {value or ''}" for key, value in row.items() if key),
                metadata={"source": name, "location": f"第 {number} 行", "row": number},
            )
            for number, row in enumerate(reader, start=2)
        ]
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")

    if not any(doc.page_content.strip() for doc in documents):
        raise ValueError(f"{name}没有可提取的文本；扫描版 PDF 需要先进行 OCR。")
    return [doc for doc in documents if doc.page_content.strip()]


def index_id(
    files: Sequence[UploadedDocument], model: str = EMBEDDING_MODEL,
    base_url: str = COMPANY_BASE_URL, embedding_provider: str = "company",
) -> str:
    """Include both filename and bytes so replacing a same-name file invalidates the cache."""
    fingerprints = sorted((display_name(f.name), hashlib.sha256(f.content).hexdigest()) for f in files)
    target = base_url if embedding_provider == "company" else LOCAL_OLLAMA_BASE_URL
    payload = repr((fingerprints, model, embedding_provider, target, CHUNK_SIZE, CHUNK_OVERLAP)).encode("utf-8")
    return "robby_" + hashlib.sha256(payload).hexdigest()[:32]



def make_embeddings(embedding_provider: str, embedding_model: str, api_key: str, base_url: str):
    """Only API embeddings receive the active service key; Ollama stays on loopback."""
    base_url = validate_base_url(base_url)
    if embedding_provider == "company":
        return OpenAIEmbeddings(
            model=embedding_model, api_key=api_key, base_url=base_url,
            chunk_size=ONLINE_BATCH_SIZE, check_embedding_ctx_length=False,
            model_kwargs={"encoding_format": "float"},
        )
    if embedding_provider == "ollama":
        try:
            from langchain_ollama import OllamaEmbeddings
        except ImportError as exc:
            raise RuntimeError(
                "本地向量化需安装 langchain-ollama："
                "pip install -r requirements-local.txt。"
            ) from exc
        return OllamaEmbeddings(model=embedding_model, base_url=LOCAL_OLLAMA_BASE_URL)
    raise ValueError("不支持的向量化方式。")


def build_index(
    files: Sequence[UploadedDocument],
    api_key: str,
    persist_directory: str | Path = "embeddings/chroma",
    embedding_model: str = EMBEDDING_MODEL,
    base_url: str = COMPANY_BASE_URL,
    embedding_provider: str = "company",
) -> IndexInfo:
    if not files:
        raise ValueError("请至少上传一个文件。")
    if embedding_provider == "company" and not api_key:
        raise ValueError("在线向量化需要提供当前模型服务的 API 密钥。")

    names = [display_name(f.name) for f in files]
    if len(names) != len(set(names)):
        raise ValueError("上传文件中有重名文件，请先重命名。")

    base_url = validate_base_url(base_url)
    collection_id = index_id(files, embedding_model, base_url, embedding_provider)
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    chunks = splitter.split_documents([doc for file in files for doc in parse_document(file)])
    chunks = [doc for doc in chunks if doc.page_content.strip()]
    ids = [hashlib.sha256(f"{collection_id}:{number}:{doc.page_content}".encode("utf-8")).hexdigest()
           for number, doc in enumerate(chunks)]

    directory = Path(persist_directory)
    directory.mkdir(parents=True, exist_ok=True)
    store = Chroma(
        collection_name=collection_id,
        persist_directory=str(directory),
        embedding_function=make_embeddings(embedding_provider, embedding_model, api_key, base_url),
    )
    existing = set()
    for start in range(0, len(ids), 256):
        existing.update(store.get(ids=ids[start:start + 256], include=[])["ids"])
    missing = [(doc, doc_id) for doc, doc_id in zip(chunks, ids) if doc_id not in existing]
    batch_size = ONLINE_BATCH_SIZE if embedding_provider == "company" else 64
    cached_count = len(ids) - len(missing)
    for start in range(0, len(missing), batch_size):
        batch = missing[start:start + batch_size]
        try:
            store.add_documents([doc for doc, _ in batch], ids=[doc_id for _, doc_id in batch])
        except Exception as exc:
            # Do not surface service error bodies: some gateways echo document text.
            status_code = getattr(exc, "status_code", None)
            raise EmbeddingBatchError(
                cached_count + start + 1, cached_count + start + len(batch), len(ids), status_code,
            ) from exc
    return IndexInfo(store, collection_id, len(files), len(chunks), not missing)
