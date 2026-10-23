"""Managed ingestion, hybrid retrieval and LangGraph tests without network calls."""

import unittest
import gc
import time
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch, MagicMock

from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_core.embeddings import FakeEmbeddings
from langchain_core.language_models.fake_chat_models import FakeListChatModel

from src.modules.chatbot import ChatResult, needs_rewrite
from src.modules.embedder import UploadedDocument, parse_document
from src.modules.repository import Repository
from src.modules.retrieval import bm25, rrf_merge
from src.modules.workflow import answer_turn, stream_answer_turn
from src.modules import managed_index


class PlatformTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.addCleanup(self.clear_chroma_cache)
        self.path = Path(self.directory.name)
        self.repo = Repository(self.path / "app.sqlite3")
        self.owner = self.repo.register("researcher", "secure-password-123")
        self.kb_id = self.repo.create_base(self.owner, "papers", "paper", "",
                                           "fake-embed", "company", "https://example.invalid/v1")
        self.kb = self.repo.get_base(self.owner, self.kb_id)

    @staticmethod
    def clear_chroma_cache():
        from chromadb.api.client import SharedSystemClient
        SharedSystemClient.clear_system_cache()
        gc.collect()

    def test_word_and_markdown_are_parsed(self):
        word = WordDocument()
        word.add_paragraph("Technical specification details")
        data = BytesIO()
        word.save(data)
        docs = parse_document(UploadedDocument("guide.docx", data.getvalue()))
        self.assertIn("Technical specification", docs[0].page_content)
        self.assertIn("paragraph", docs[0].metadata)
        markdown = parse_document(UploadedDocument("README.md", b"# Title\nImportant detail"))
        self.assertIn("Important detail", markdown[0].page_content)

    def test_upload_async_then_hybrid_retrieval_without_network(self):
        from src.modules.retrieval import HybridRetriever
        with patch.object(managed_index, "DATA_DIR", self.path), \
                patch.object(managed_index, "make_embeddings", return_value=FakeEmbeddings(size=8)), \
                patch("src.modules.retrieval.vector_store", side_effect=managed_index.vector_store):
            doc, created = managed_index.queue_upload(self.repo, self.owner, self.kb,
                                                       UploadedDocument("demo.md", b"# RAG\nHybrid retrieval with evidence"), "fake-key")
            self.assertTrue(created)
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline and self.repo.get_document(self.owner, doc["id"])["status"] != "success":
                time.sleep(.05)
            self.assertEqual(self.repo.get_document(self.owner, doc["id"])["status"], "success")
            retriever = HybridRetriever(self.repo, self.owner, self.kb, "fake-key", "fake-model", mode="hybrid")
            found = retriever.similarity_search("hybrid retrieval", k=4)
            self.assertTrue(found)
            self.assertEqual(found[0].metadata["source"], "demo.md")
            self.assertEqual(found[0].metadata["document_id"], doc["id"])
            while time.monotonic() < deadline and doc["id"] in managed_index._pending:
                time.sleep(.05)

    def test_async_job_body_and_versioned_vectors_are_real(self):
        with patch.object(managed_index, "DATA_DIR", self.path), \
                patch.object(managed_index, "make_embeddings", return_value=FakeEmbeddings(size=8)), \
                patch.object(managed_index, "enqueue", return_value=True):
            first, created = managed_index.queue_upload(self.repo, self.owner, self.kb,
                                                         UploadedDocument("paper.txt", b"research evidence alpha"), "fake-key")
            self.assertTrue(created)
            managed_index._process(self.repo.path, self.owner, self.kb_id, first["id"], "fake-key")
            self.assertEqual(self.repo.get_document(self.owner, first["id"])["status"], "success")
            initial_ids = [r["id"] for r in self.repo.list_chunks(self.owner, self.kb_id)]
            second, created = managed_index.queue_upload(self.repo, self.owner, self.kb,
                                                          UploadedDocument("paper.txt", b"research evidence beta"), "fake-key")
            self.assertTrue(created)
            self.assertEqual(second["version"], 2)
            managed_index._process(self.repo.path, self.owner, self.kb_id, second["id"], "fake-key")
            active = self.repo.list_chunks(self.owner, self.kb_id)
            self.assertNotEqual(active[0]["id"], initial_ids[0])
            self.assertEqual(self.repo.stats(self.owner)["documents"], 1)
            self.assertFalse(self.repo.get_document(self.owner, first["id"])["active"])
            from src.modules.retrieval import HybridRetriever
            with patch("src.modules.retrieval.vector_store", side_effect=lambda kb, key: managed_index.vector_store(kb, key)):
                retriever = HybridRetriever(self.repo, self.owner, self.kb, "fake-key", "fake-model")
                results = retriever.similarity_search("research evidence", k=2)
                self.assertTrue(results)
                self.assertEqual(results[0].metadata["document_id"], second["id"])
            managed_index.delete_document(self.repo, self.owner, self.kb, second["id"], "fake-key")
            self.assertEqual(self.repo.stats(self.owner)["documents"], 0)
            self.assertEqual(self.repo.list_chunks(self.owner, self.kb_id), [])

    def test_bm25_rrf_and_topic_switch(self):
        rows = [
            {"id": "a", "content": "Transformer attention sequence", "source": "a.txt", "location": "1", "document_id": "d", "version": 1},
            {"id": "b", "content": "company annual leave 年假 政策", "source": "b.txt", "location": "2", "document_id": "e", "version": 1},
        ]
        ranked = bm25("公司的年假有几天", rows, 2)
        self.assertEqual(ranked[0].metadata["chunk_id"], "b")
        other = Document(page_content=rows[0]["content"], metadata={"chunk_id": "a"})
        fused = rrf_merge([[other], ranked], 2)
        self.assertEqual(len(fused), 2)
        self.assertFalse(needs_rewrite("公司的年假有几天？"))
        self.assertTrue(needs_rewrite("那它有什么缺点？"))

    def test_listwise_rerank_uses_real_model_result_order(self):
        from src.modules.retrieval import llm_rerank
        docs = [Document(page_content=f"Evidence {i}", metadata={"chunk_id": str(i)}) for i in range(1, 4)]
        with patch("src.modules.retrieval.ChatOpenAI") as model:
            model.return_value.invoke.return_value.content = "[3,1,2]"
            ranked = llm_rerank("research question", docs, "test-model", "test-key", "https://example.invalid/v1")
        self.assertEqual([d.metadata["chunk_id"] for d in ranked], ["3", "1", "2"])
        model.assert_called_once()

    def test_stopped_stream_does_not_persist_partial_answer(self):
        sid = self.repo.create_session(self.owner, self.kb_id)
        evidence = Document(page_content="evidence", metadata={"chunk_id": "c1", "source": "paper.pdf"})
        with patch("src.modules.workflow.HybridRetriever") as retriever, \
                patch("src.modules.chatbot.ChatOpenAI", return_value=FakeListChatModel(responses=["Long answer [1]"])):
            retriever.return_value.similarity_search.return_value = [evidence]
            stream = stream_answer_turn(self.repo, self.owner, self.kb, sid,
                                        "Question?", "fake-key", "fake-model")
            self.assertTrue(next(stream))
            stream.close()
        self.assertEqual(self.repo.list_messages(self.owner, sid), [])

    def test_stream_turn_yields_real_model_chunks(self):
        sid = self.repo.create_session(self.owner, self.kb_id)
        evidence = Document(page_content="research evidence", metadata={"chunk_id": "c1", "source": "paper.pdf"})
        with patch("src.modules.workflow.HybridRetriever") as retriever, \
                patch("src.modules.chatbot.ChatOpenAI", return_value=FakeListChatModel(responses=["Answer [1]"])):
            retriever.return_value.similarity_search.return_value = [evidence]
            chunks = list(stream_answer_turn(self.repo, self.owner, self.kb, sid,
                                             "Question?", "fake-key", "fake-model"))
        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(chunks), "Answer [1]")
        self.assertEqual(self.repo.list_messages(self.owner, sid)[-1]["content"], "Answer [1]")

    def test_stream_turn_commits_only_on_complete(self):
        sid = self.repo.create_session(self.owner, self.kb_id)
        evidence = Document(page_content="evidence", metadata={"chunk_id": "a", "source": "a.txt"})
        fake_bot = MagicMock()
        fake_bot.rewrite_question.return_value = "question"
        fake_bot.retrieve.return_value = [evidence]
        fake_bot.answer.return_value = ChatResult("Evidence [1]", [evidence])
        with patch("src.modules.workflow.HybridRetriever"), patch("src.modules.workflow.Chatbot", return_value=fake_bot):
            response = stream_answer_turn(self.repo, self.owner, self.kb, sid,
                                          "question", "fake-key", "fake-model")
            self.assertEqual(self.repo.list_messages(self.owner, sid), [])
            self.assertEqual("".join(response), "Evidence [1]")
        messages = self.repo.list_messages(self.owner, sid)
        self.assertEqual([item["role"] for item in messages], ["user", "assistant"])
        self.assertEqual(messages[-1]["sources"][0]["name"], "a.txt")

    def test_regenerate_replaces_answer_not_history(self):
        sid = self.repo.create_session(self.owner, self.kb_id)
        self.repo.append_turn(self.owner, sid, "question", "old answer", [])
        fake_bot = MagicMock()
        fake_bot.rewrite_question.return_value = "question"
        fake_bot.retrieve.return_value = []
        fake_bot.answer.return_value = ChatResult("new answer", [])
        with patch("src.modules.workflow.HybridRetriever"), patch("src.modules.workflow.Chatbot", return_value=fake_bot):
            answer_turn(self.repo, self.owner, self.kb, sid, "question", "fake-key", "fake-model", regenerate=True)
        messages = self.repo.list_messages(self.owner, sid)
        self.assertEqual(len(messages), 2)
        self.assertEqual(messages[-1]["content"], "new answer")
    def test_graph_persists_grounded_turn_without_cross_session(self):
        sid = self.repo.create_session(self.owner, self.kb_id)
        doc = Document(page_content="research evidence", metadata={"source": "paper.pdf", "location": "page 1", "chunk_id": "c1"})
        fake_bot = MagicMock()
        fake_bot.rewrite_question.return_value = "independent query"
        fake_bot.retrieve.return_value = [doc]
        fake_bot.answer.return_value = ChatResult("Answer [1]", [doc])
        with patch("src.modules.workflow.HybridRetriever"), patch("src.modules.workflow.Chatbot", return_value=fake_bot):
            result, metrics = answer_turn(self.repo, self.owner, self.kb, sid,
                                          "question", "fake-key", "fake-model")
        self.assertEqual(result.answer, "Answer [1]")
        messages = Repository(self.repo.path).list_messages(self.owner, sid)
        self.assertEqual([m["role"] for m in messages], ["user", "assistant"])
        self.assertEqual(messages[1]["sources"][0]["location"], "page 1")
        other = self.repo.register("otheruser", "another-password-123")
        with self.assertRaises(PermissionError):
            answer_turn(self.repo, other, self.kb, sid, "question", "fake-key", "fake-model")


if __name__ == "__main__":
    unittest.main()